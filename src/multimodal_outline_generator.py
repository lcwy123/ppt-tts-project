#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
多模态提纲生成模块 - 从Word文档提取的内容生成PPT提纲
基于标题层级组织，每级标题一页PPT
"""

import json
import re
from pathlib import Path
from typing import Dict, List, Any, Tuple

from src.config import get_llm_client, get_llm_model_name, OUTPUT_OUTLINE, LLM_MODE


def run(extracted_data: Dict[str, Any], progress_callback=None) -> Tuple[str, str]:
    """
    根据提取的文档内容生成PPT提纲

    Args:
        extracted_data: docx_extractor.extract_from_docx() 返回的数据
        progress_callback: 进度回调函数

    Returns:
        (提纲Word文件路径, 提纲文本)
    """
    if progress_callback:
        progress_callback("开始生成多模态提纲...")

    # 直接基于提取的内容块生成提纲
    # 每个heading级别的块就是一页PPT
    slides_data = []

    # 生成封面页
    title = extracted_data.get("title", "PPT提纲")
    slides_data.append({
        "page_num": 1,
        "title": title,
        "content": "封面",
        "images": [],
        "is_cover": True
    })

    # 遍历内容块，每块一页
    page_num = 2
    for block in extracted_data.get("slides", []):
        block_title = block.get("title", "")
        block_text = block.get("text", "")
        block_images = block.get("images", [])
        block_level = block.get("level", 0)

        if not block_title:
            continue

        # 跳过纯标题（没有实质内容）
        if not block_text or block_text == block_title:
            continue

        slides_data.append({
            "page_num": page_num,
            "title": block_title,
            "content": block_text,
            "images": block_images,  # 直接使用该标题关联的图片
            "is_cover": False,
            "level": block_level
        })
        page_num += 1

    # 添加结束页
    slides_data.append({
        "page_num": page_num,
        "title": "谢谢观看",
        "content": "",
        "images": [],
        "is_cover": False
    })

    if progress_callback:
        progress_callback(f"提纲生成成功，共{len(slides_data)}页")

    # 生成提纲文本
    outline_text = generate_outline_text(slides_data)
    outline_path = create_word_outline(slides_data)

    return outline_path, outline_text


def generate_outline_text(slides_data: List[Dict]) -> str:
    """生成提纲文本"""
    lines = []
    lines.append("PPT提纲")
    lines.append("")

    for slide in slides_data:
        page_num = slide.get("page_num", 0)
        title = slide.get("title", "")
        content = slide.get("content", "")
        images = slide.get("images", [])
        is_cover = slide.get("is_cover", False)

        lines.append(f"=== 第{page_num}页 ===")
        lines.append(f"# {title}")
        lines.append("")

        if is_cover:
            lines.append(content)
        else:
            # 内容按行分割
            for line in content.split("\n"):
                line = line.strip()
                if line:
                    lines.append(f"- {line}")

            # 标注关联图片
            if images:
                lines.append("")
                lines.append(f"📷 配图: {', '.join(images)}")

        lines.append("")

    return "\n".join(lines)


def create_word_outline(slides_data: List[Dict]) -> str:
    """创建Word提纲文件"""
    from docx import Document
    from docx.shared import Pt
    from docx.enum.text import WD_PARAGRAPH_ALIGNMENT

    doc = Document()

    # 设置标题
    title = doc.add_heading("PPT提纲", 0)
    title.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

    for slide in slides_data:
        page_num = slide.get("page_num", 0)
        title_text = slide.get("title", "")
        content = slide.get("content", "")
        images = slide.get("images", [])

        # 页面标题
        heading = doc.add_heading(f"{title_text}", level=1)

        # 内容
        for line in content.split("\n"):
            line = line.strip()
            if line:
                doc.add_paragraph(line, style='List Bullet')

        # 添加图片占位符（如果有）
        if images:
            p = doc.add_paragraph()
            run = p.add_run("📷 配图: ")
            run.bold = True
            run = p.add_run(", ".join(images))

        doc.add_paragraph()  # 空行

    # 保存
    output_path = OUTPUT_OUTLINE
    output_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(output_path))

    return str(output_path)
