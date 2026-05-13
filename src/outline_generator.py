#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
文字转Word提纲模块
"""

import os
from pathlib import Path
from docx import Document
from docx.shared import Pt
from src.config import (
    get_llm_client, get_llm_model_name, OUTPUT_OUTLINE, INPUT_DIR, LLM_MODE
)

def generate_outline_from_text(input_text, progress_callback=None):
    """使用LLM从输入文字生成PPT提纲"""

    client = get_llm_client()
    model = get_llm_model_name()
    
    prompt = f"""
请为以下主题内容生成一个完整的PPT提纲，格式要求如下：

1. 第一行是PPT标题
2. 然后按幻灯片分，每一页幻灯片用"=== 第X页 ==="分隔
3. 每页包含：页面标题（用#开头），要点列表（用-开头）

示例格式：
人工智能发展概述
=== 第1页 ===
# 目录
- 背景介绍
- 发展历程
- 核心技术
- 未来展望

=== 第2页 ===
# 背景介绍
- 人工智能的定义
- 人工智能的发展历史
- 当前技术现状

请为以下内容生成提纲：

{input_text}
"""
    
    if progress_callback:
        progress_callback(f"正在调用LLM生成提纲...")
    
    try:
        if LLM_MODE == "api":
            # Anthropic/MiniMax API
            response = client.messages.create(
                model=model,
                max_tokens=4096,
                system="You are a helpful assistant.",
                messages=[
                    {"role": "user", "content": prompt}
                ]
            )
            # 遍历content获取文本，跳过thinking块
            outline_text = ""
            for block in response.content:
                if block.type == 'text':
                    outline_text = block.text
                    break
        else:
            # OpenAI/SGlang API
            response = client.chat.completions.create(
                model=model,
                messages=[{"role": "user", "content": prompt}],
                temperature=0.7,
            )
            outline_text = response.choices[0].message.content

        if progress_callback:
            progress_callback(f"提纲生成成功")

        return outline_text

    except Exception as e:
        if progress_callback:
            progress_callback(f"LLM调用失败: {str(e)}，使用离线模式")
        return generate_outline_offline(input_text)


def generate_outline_offline(input_text):
    """离线生成提纲（基于模板）"""
    
    # 简单模板生成
    title = input_text.split('\n')[0][:50] if input_text else "PPT主题"
    
    outline = f"""{title}

=== 第1页 ===
# 目录
- 第一部分内容
- 第二部分内容
- 第三部分内容

=== 第2页 ===
# 第一部分
- 要点1
- 要点2
- 要点3

=== 第3页 ===
# 第二部分
- 要点1
- 要点2
- 要点3

=== 第4页 ===
# 第三部分
- 要点1
- 要点2
- 要点3

=== 第5页 ===
# 总结
- 核心要点回顾
- 谢谢观看
"""
    return outline


def create_word_outline(outline_text, output_path=None, progress_callback=None):
    """将提纲文本生成Word文档"""
    
    if output_path is None:
        output_path = OUTPUT_OUTLINE
    
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    
    doc = Document()
    style = doc.styles['Normal']
    font = style.font
    font.name = '宋体'
    font.size = Pt(12)
    
    lines = outline_text.strip().split('\n')
    first_line = True
    
    for line in lines:
        line = line.strip()
        if not line:
            continue
            
        if first_line:
            heading = doc.add_heading(line, level=0)
            heading.runs[0].font.name = '黑体'
            first_line = False
        elif line.startswith('==='):
            doc.add_page_break()
        elif line.startswith('# '):
            heading = doc.add_heading(line[2:], level=1)
            heading.runs[0].font.name = '黑体'
        elif line.startswith('- '):
            p = doc.add_paragraph(line[2:], style='List Bullet')
            p.runs[0].font.name = '宋体'
        else:
            p = doc.add_paragraph(line)
            p.runs[0].font.name = '宋体'
    
    doc.save(output_path)
    
    if progress_callback:
        progress_callback(f"Word提纲已保存: {output_path}")
    
    return output_path


def run(input_text=None, progress_callback=None):
    """运行完整流程：文字→提纲→Word"""
    
    if input_text is None:
        input_file = INPUT_DIR / "input.txt"
        if input_file.exists():
            with open(input_file, 'r', encoding='utf-8') as f:
                input_text = f.read()
        else:
            raise FileNotFoundError(f"输入文件不存在: {input_file}")
    
    if progress_callback:
        progress_callback("开始生成提纲...")
    
    outline_text = generate_outline_from_text(input_text, progress_callback)
    output_path = create_word_outline(outline_text, progress_callback=progress_callback)
    
    return output_path, outline_text
