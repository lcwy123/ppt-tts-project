#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
音频嵌入PPT模块
"""

import os
import json
from pathlib import Path
from pptx import Presentation
from pptx.util import Inches
from src.config import OUTPUT_PPT, INPUT_DIR, OUTPUT_DIR, OUTPUT_AUDIO_DURATIONS, OUTPUT_FINAL_PPT


def load_audio_durations(file_path):
    """加载音频时长信息"""
    
    with open(file_path, 'r', encoding='utf-8') as f:
        return json.load(f)


def embed_audio_to_slide(slide, audio_path, duration, slide_index):
    """将音频嵌入单页幻灯片（放置在可视区域外隐藏）"""
    
    # 添加音频（放置在幻灯片外，隐藏图标）
    left = Inches(-1)
    top = Inches(-1)
    width = Inches(0.5)
    height = Inches(0.5)
    
    audio_shape = slide.shapes.add_movie(
        audio_path,
        left, top, width, height,
        poster_frame_image=None,
        mime_type='audio/mp3'
    )
    
    return audio_shape


def run(ppt_path=None, durations_path=None, progress_callback=None):
    """运行完整流程：嵌入音频到PPT"""
    
    # 确定PPT文件
    if ppt_path is None:
        input_ppt = INPUT_DIR / "input.pptx"
        generated_ppt = OUTPUT_PPT
        
        if input_ppt.exists():
            ppt_path = input_ppt
        elif generated_ppt.exists():
            ppt_path = generated_ppt
        else:
            raise FileNotFoundError("未找到PPT文件")
    
    # 确定时长文件
    if durations_path is None:
        durations_path = OUTPUT_AUDIO_DURATIONS
    
    if not os.path.exists(durations_path):
        raise FileNotFoundError(f"音频时长文件不存在: {durations_path}")
    
    if progress_callback:
        progress_callback("正在加载音频时长信息...")
    
    audio_durations = load_audio_durations(durations_path)
    audio_map = {item["slide_num"]: item for item in audio_durations}
    
    # 打开PPT
    prs = Presentation(ppt_path)
    total_slides = len(prs.slides)
    
    if progress_callback:
        progress_callback(f"PPT共 {total_slides} 页幻灯片")
    
    embedded_count = 0
    
    for i, slide in enumerate(prs.slides):
        slide_num = i + 1
        
        if slide_num in audio_map:
            audio_info = audio_map[slide_num]
            audio_path = audio_info["audio_path"]
            duration = audio_info["duration"]
            
            if os.path.exists(audio_path):
                if progress_callback:
                    progress_callback(f"正在嵌入第 {slide_num} 页音频 (时长: {duration:.1f}秒)...")
                embed_audio_to_slide(slide, audio_path, duration, i)
                embedded_count += 1
            else:
                if progress_callback:
                    progress_callback(f"第 {slide_num} 页音频文件不存在: {audio_path}")
        else:
            if progress_callback:
                progress_callback(f"第 {slide_num} 页没有音频")
    
    # 保存最终PPT
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    output_path = OUTPUT_FINAL_PPT
    prs.save(output_path)
    
    if progress_callback:
        progress_callback(f"\n✅ 音频嵌入完成！成功嵌入 {embedded_count} 页")
        progress_callback(f"📁 最终PPT: {output_path}")
    
    return output_path, embedded_count
