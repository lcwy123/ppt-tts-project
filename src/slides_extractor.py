#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
PPT文本提取模块
"""

import os
from pathlib import Path
from pptx import Presentation
from pptx.shapes.group import GroupShape
from src.config import OUTPUT_PPT, INPUT_DIR, OUTPUT_SLIDES_TEXT


def get_all_text_from_shape(shape):
    """递归获取形状及其子形状的所有文本"""
    texts = []
    
    # 获取当前形状的文本
    if hasattr(shape, "text") and shape.text.strip():
        texts.append(shape.text.strip())
    
    # 如果是组合形状，递归获取子形状文本
    if isinstance(shape, GroupShape):
        for sub_shape in shape.shapes:
            texts.extend(get_all_text_from_shape(sub_shape))
    
    return texts


def extract_slides_text(ppt_path, progress_callback=None):
    """提取PPT每页的文本内容"""
    
    prs = Presentation(ppt_path)
    slides_text = []
    
    for i, slide in enumerate(prs.slides):
        slide_content = []
        
        for shape in slide.shapes:
            slide_content.extend(get_all_text_from_shape(shape))
        
        # 去重并保持顺序
        seen = set()
        unique_content = []
        for text in slide_content:
            if text not in seen:
                seen.add(text)
                unique_content.append(text)
        
        slides_text.append({
            "slide_num": i + 1,
            "content": "\n".join(unique_content)
        })
        
        if progress_callback:
            progress_callback(f"提取第 {i+1}/{len(prs.slides)} 页...")
    
    return slides_text


def save_slides_text(slides_text, output_path=None, progress_callback=None):
    """保存提取的文本到文件"""
    
    if output_path is None:
        output_path = OUTPUT_SLIDES_TEXT
    
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    
    with open(output_path, 'w', encoding='utf-8') as f:
        for slide in slides_text:
            f.write(f"=== 第 {slide['slide_num']} 页 ===\n")
            f.write(slide["content"])
            f.write("\n\n")
    
    if progress_callback:
        progress_callback(f"已提取 {len(slides_text)} 页文本，保存到: {output_path}")
    
    return output_path


def run(ppt_path=None, progress_callback=None):
    """运行完整流程：提取PPT文本"""
    
    if ppt_path is None:
        # 优先使用生成的ppt，其次是input目录的ppt
        generated_ppt = OUTPUT_PPT
        input_ppt = INPUT_DIR / "input.pptx"
        
        # 优先使用生成的PPT（如果存在且非空）
        if generated_ppt.exists() and generated_ppt.stat().st_size > 0:
            ppt_path = generated_ppt
            if progress_callback:
                progress_callback(f"使用生成的PPT: {ppt_path}")
        elif input_ppt.exists() and input_ppt.stat().st_size > 0:
            ppt_path = input_ppt
            if progress_callback:
                progress_callback(f"使用输入PPT: {ppt_path}")
        else:
            raise FileNotFoundError("未找到PPT文件")
    
    if progress_callback:
        progress_callback("正在提取PPT文本...")
    
    slides_text = extract_slides_text(ppt_path, progress_callback)
    output_path = save_slides_text(slides_text, progress_callback=progress_callback)
    
    return output_path, slides_text
